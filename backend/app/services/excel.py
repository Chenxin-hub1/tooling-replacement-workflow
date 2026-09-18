import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Literal
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.dates import local_today
from app.engine.scheduler import (
    build_actions_for_project,
    compute_project_summary,
    reassign_team_member,
    schedule_project_actions,
)
from app.models import Project, ProjectAction, TeamMember
from app.schemas.excel import (
    ImportAction,
    ImportIssue,
    ImportProject,
    ImportReport,
    ImportTeam,
)
from app.services.projects import touch

Mode = Literal["create", "upsert"]
MAX_UPLOAD = 10 * 1024 * 1024
MAX_ROWS = 10000
PROJECT_FIELDS = list(ImportProject.model_fields)
PROJECT_FIELDS.remove("team")
TEAM_FIELDS = ["project_id", "function", "name", "email"]
ACTION_FIELDS = [
    "id",
    "project_id",
    *[key for key in ImportAction.model_fields if key not in ("id", "project_id")],
]
ALIASES = {
    "id": [
        "project id",
        "project number",
        "项目编号",
        "项目号",
        "action id",
        "动作编号",
    ],
    "project_id": ["project_id", "project", "项目", "项目编号"],
    "pn": ["part number", "part no", "零件号", "料号"],
    "desc": ["description", "part description", "零件描述", "描述"],
    "plant": ["plant affected", "工厂"],
    "bu": ["business unit", "业务单元"],
    "cur": ["current supplier", "old supplier", "旧供应商"],
    "nw": ["new supplier", "新供应商"],
    "po": ["tool po", "purchase order", "采购订单"],
    "tag": ["tool tag", "asset number", "资产编号"],
    "too_owner": ["tool owner", "模具归属"],
    "cav": ["cavities", "模穴数"],
    "saving": ["annual saving", "estimated saving", "年度节省"],
    "oem": ["customer", "客户"],
    "tech": ["technology", "工艺"],
    "reason": ["replacement reason", "原因"],
    "owner": ["assignee", "responsible", "负责人"],
    "created_at": ["created date", "立项日期"],
    "function": ["role", "职能", "角色"],
    "name": ["member", "姓名"],
    "email": ["邮箱"],
    "ph": ["phase", "阶段"],
    "tab": ["group", "分组"],
    "act": ["action", "动作"],
    "input_type": ["input type", "输入类型"],
    "fn": ["role", "function", "职能", "角色"],
    "dep_id": ["predecessor", "dependency", "前置动作"],
    "lead": ["lead time", "交期"],
    "due_date": ["due date", "target date", "目标日期"],
    "done_date": ["done date", "completion date", "完成日期"],
    "value": ["input value", "填报值"],
    "comment": ["comments", "备注"],
    "link": ["document link", "链接"],
    "is_custom": ["custom", "自定义"],
}


def normalize(value: object) -> str:
    return re.sub(r"[\s_\-./()]+", "", str(value or "")).lower()


@dataclass
class WorkbookData:
    projects: list[tuple[int, ImportProject]] = field(default_factory=list)
    teams: list[tuple[int, ImportTeam]] = field(default_factory=list)
    actions: list[tuple[int, ImportAction]] = field(default_factory=list)
    errors: list[ImportIssue] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def parse_workbook(content: bytes) -> WorkbookData:
    result = WorkbookData()
    if len(content) > MAX_UPLOAD:
        raise ValueError("Workbook exceeds the 10 MB upload limit")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Workbook expands beyond the 50 MB limit")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Upload a valid .xlsx workbook") from exc
    definitions = {
        "Projects": (
            ImportProject,
            result.projects,
            {"id", "pn", "desc"},
            ["projects", "master changes", "master", "项目"],
        ),
        "Teams": (
            ImportTeam,
            result.teams,
            {"project_id", "function", "name"},
            ["teams", "team", "团队"],
        ),
        "Actions": (
            ImportAction,
            result.actions,
            {"id", "project_id", "ph", "tab", "act", "input_type", "fn"},
            ["actions", "oil", "动作"],
        ),
    }
    try:
        for canonical, (model, output, required, sheet_names) in definitions.items():
            sheets = [
                sheet
                for sheet in book
                if normalize(sheet.title) in {normalize(name) for name in sheet_names}
            ]
            if not sheets:
                if canonical == "Projects":
                    result.errors.append(
                        ImportIssue(
                            sheet=canonical,
                            row=0,
                            message="Missing Projects (or Master Changes) sheet",
                        )
                    )
                continue
            if len(sheets) > 1:
                result.errors.append(
                    ImportIssue(
                        sheet=canonical,
                        row=0,
                        message="Multiple sheets map to the same dataset",
                    )
                )
                continue
            sheet = sheets[0]
            if sheet.max_column and sheet.max_column > 1000:
                result.errors.append(
                    ImportIssue(
                        sheet=sheet.title, row=0, message="Sheet exceeds 1000 columns"
                    )
                )
                continue
            lookup = {}
            for key in model.model_fields:
                for alias in [key, *ALIASES.get(key, [])]:
                    lookup[normalize(alias)] = key
            # 项目编号在主表表示 id，在动作与团队表表示 project_id。
            lookup[normalize("项目编号")] = (
                "id" if canonical == "Projects" else "project_id"
            )
            rows = sheet.iter_rows()
            mapping = None
            header_row = 0
            for number, cells in enumerate(rows, 1):
                recognized = {
                    index: lookup[normalize(cell.value)]
                    for index, cell in enumerate(cells)
                    if normalize(cell.value) in lookup
                }
                if required <= set(recognized.values()):
                    mapping, header_row = recognized, number
                    if len(set(mapping.values())) != len(mapping):
                        result.errors.append(
                            ImportIssue(
                                sheet=sheet.title,
                                row=number,
                                message="Duplicate mapped columns",
                            )
                        )
                    unknown = [
                        str(cell.value)
                        for index, cell in enumerate(cells)
                        if cell.value is not None and index not in recognized
                    ]
                    if unknown:
                        result.warnings.append(
                            f"{sheet.title}: ignored columns: {', '.join(unknown)}"
                        )
                    break
                if number >= 20:
                    break
            if mapping is None:
                result.errors.append(
                    ImportIssue(
                        sheet=sheet.title,
                        row=0,
                        message=f"Missing required headers: {', '.join(sorted(required))}",
                    )
                )
                continue
            for number, cells in enumerate(rows, header_row + 1):
                if number > MAX_ROWS + header_row:
                    result.errors.append(
                        ImportIssue(
                            sheet=sheet.title,
                            row=number,
                            message="Sheet exceeds 10000 data rows",
                        )
                    )
                    break
                if not any(cell.value is not None for cell in cells):
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    result.errors.append(
                        ImportIssue(
                            sheet=sheet.title,
                            row=number,
                            message="Formula cells are not supported; paste values before importing",
                        )
                    )
                    continue
                values = {}
                for index, key in mapping.items():
                    value = cells[index].value if index < len(cells) else None
                    explicit_empty = (
                        index < len(cells)
                        and cells[index].data_type == "inlineStr"
                        and value is None
                    )
                    if explicit_empty and key not in (
                        "ph",
                        "lead",
                        "lead_default",
                        "is_custom",
                        "created_at",
                        "due_date",
                        "done_date",
                        "last_remind_date",
                    ):
                        value = ""
                    if isinstance(value, datetime):
                        value = value.date().isoformat()
                    elif isinstance(value, date):
                        value = value.isoformat()
                    if value == "" and not explicit_empty:
                        value = None
                    if (
                        key not in ("ph", "lead", "lead_default", "is_custom")
                        and value is not None
                    ):
                        value = str(value)
                    if key == "is_custom" and value is None:
                        value = False
                    values[key] = value
                try:
                    item = model.model_validate(values)
                    if isinstance(item, ImportProject):
                        result.projects.append((number, item))
                    elif isinstance(item, ImportTeam):
                        result.teams.append((number, item))
                    else:
                        result.actions.append((number, item))
                except ValidationError as exc:
                    message = "; ".join(
                        f"{'.'.join(map(str, error['loc']))}: {error['msg']}"
                        for error in exc.errors()
                    )
                    result.errors.append(
                        ImportIssue(sheet=sheet.title, row=number, message=message)
                    )
    finally:
        book.close()
    return result


def clone_project(project: Project) -> Project:
    clone = Project(
        **{
            column.key: getattr(project, column.key)
            for column in Project.__table__.columns
        }
    )
    clone.team_members = [
        TeamMember(
            **{
                column.key: getattr(member, column.key)
                for column in TeamMember.__table__.columns
            }
        )
        for member in project.team_members
    ]
    clone.actions = [
        ProjectAction(
            **{
                column.key: getattr(action, column.key)
                for column in ProjectAction.__table__.columns
            }
        )
        for action in project.actions
    ]
    return clone


async def prepare_import(
    db: AsyncSession, data: WorkbookData, mode: Mode
) -> tuple[ImportReport, list[Project]]:
    report = ImportReport(
        valid=False,
        errors=list(data.errors),
        warnings=list(data.warnings),
        action_rows=len(data.actions),
    )
    planned: dict[str, Project] = {}

    def error(sheet: str, row: int, message: str):
        report.errors.append(ImportIssue(sheet=sheet, row=row, message=message))

    for row, item in data.projects:
        if item.id in planned:
            error("Projects", row, f"Duplicate project ID: {item.id}")
            continue
        existing = await db.get(Project, item.id)
        if existing is not None:
            if mode == "create":
                error(
                    "Projects",
                    row,
                    f"Project {item.id} already exists; choose update mode",
                )
            project = clone_project(existing)
            for key, value in item.model_dump(
                exclude_unset=True, exclude={"team"}
            ).items():
                if key != "created_at" or value is not None:
                    setattr(project, key, value)
            report.updated_projects += 1
        else:
            values = item.model_dump(exclude={"team", "created_at"})
            project = Project(**values, created_at=item.created_at or local_today())
            project.team_members = []
            project.actions = build_actions_for_project(item.id, {})
            report.created_projects += 1
        planned[item.id] = project
    seen_teams = set()
    for row, item in data.teams:
        project = planned.get(item.project_id)
        if project is None:
            error("Teams", row, "Project must be included in Projects sheet")
            continue
        key = (item.project_id, item.function)
        if key in seen_teams:
            error("Teams", row, "Duplicate role for project")
            continue
        seen_teams.add(key)
        member = next(
            (m for m in project.team_members if m.function == item.function), None
        )
        if member is None:
            member = TeamMember(
                project_id=item.project_id, function=item.function, name=item.name
            )
            project.team_members.append(member)
        member.name = item.name
        if "email" in item.model_fields_set:
            member.email = item.email
        reassign_team_member(project.actions, item.function, item.name)
    seen_actions = set()
    global_ids = {
        identifier: project_id
        for identifier, project_id in (
            await db.execute(select(ProjectAction.id, ProjectAction.project_id))
        ).all()
    }
    for row, item in data.actions:
        project = planned.get(item.project_id)
        if project is None:
            error("Actions", row, "Project must be included in Projects sheet")
            continue
        if item.id in seen_actions or (
            item.id in global_ids and global_ids[item.id] != item.project_id
        ):
            error("Actions", row, f"Duplicate or cross-project action ID: {item.id}")
            continue
        seen_actions.add(item.id)
        action = next((a for a in project.actions if a.id == item.id), None)
        values = item.model_dump(exclude_unset=True)
        if action is None:
            if not item.is_custom:
                error("Actions", row, "Unknown standard action ID")
                continue
            action = ProjectAction(**item.model_dump())
            project.actions.append(action)
        elif action.is_custom != item.is_custom:
            error(
                "Actions", row, "Cannot change whether an action is standard or custom"
            )
            continue
        elif not action.is_custom:
            immutable = (
                "ph",
                "tab",
                "act",
                "input_type",
                "fn",
                "dep_id",
                "lead_default",
            )
            if any(
                key in values and values[key] != getattr(action, key)
                for key in immutable
            ):
                error(
                    "Actions",
                    row,
                    "Standard action definition differs from the current matrix",
                )
                continue
        for key, value in values.items():
            setattr(action, key, value)
    for project in planned.values():
        try:
            schedule_project_actions(project.created_at, project.actions)
            touch(project)
        except ValueError as exc:
            error("Actions", 0, f"{project.id}: {exc}")
    report.project_ids = list(planned)
    report.valid = not report.errors
    return report, list(planned.values())


def export_workbook(projects: list[Project]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    tables = {
        "Projects": (
            PROJECT_FIELDS,
            [[getattr(p, key) for key in PROJECT_FIELDS] for p in projects],
        ),
        "Teams": (
            TEAM_FIELDS,
            [
                [getattr(member, key) for key in TEAM_FIELDS]
                for p in projects
                for member in p.team_members
            ],
        ),
        "Actions": (
            ACTION_FIELDS,
            [
                [getattr(action, key) for key in ACTION_FIELDS]
                for p in projects
                for action in p.actions
            ],
        ),
        "Summary": (
            ["project_id", "phase", "health", "progress", "overdue", "due_soon"],
            [
                [
                    p.id,
                    (s := compute_project_summary(p.created_at, p.actions)).phase,
                    s.overall_status,
                    s.pct,
                    s.overdue_actions,
                    s.due_soon_actions,
                ]
                for p in projects
            ],
        ),
    }
    for title, (headers, rows) in tables.items():
        sheet = book.create_sheet(title)
        sheet.append(headers)
        for values in rows:
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="153D3A")
            sheet.column_dimensions[cell.column_letter].width = 24
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
