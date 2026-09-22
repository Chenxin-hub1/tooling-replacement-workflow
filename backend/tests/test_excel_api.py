import json
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, text
from sqlalchemy.ext.asyncio import async_sessionmaker

from app.engine.scheduler import STANDARD_MATRIX
from app.models import Project


def workbook(rows, title="Projects", teams=None, actions=None):
    book = Workbook()
    sheet = book.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    if teams:
        sheet = book.create_sheet("Teams")
        for row in teams:
            sheet.append(row)
    if actions:
        sheet = book.create_sheet("Actions")
        for row in actions:
            sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def upload(content):
    return {
        "file": (
            "test.xlsx",
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }


async def test_preview_is_read_only_then_import_persists(client):
    content = workbook(
        [["id", "pn", "desc"], ["XL-1", "00123", "Housing"]],
        teams=[["project_id", "function", "name"], ["XL-1", "SDE", "Alice"]],
    )
    response = await client.post("/api/excel/preview", files=upload(content))
    assert response.status_code == 200, response.text
    assert response.json()["valid"]
    assert response.json()["created_projects"] == 1
    assert (await client.get("/api/projects")).json() == []
    response = await client.post("/api/excel/import", files=upload(content))
    assert response.status_code == 200, response.text
    project = (await client.get("/api/projects/XL-1")).json()
    assert project["pn"] == "00123"
    assert len(project["actions"]) == 20
    assert project["actions"][0]["owner"] == "Alice"


async def test_import_aliases_header_offset_and_excel_date(client):
    content = workbook(
        [
            ["Tooling master list"],
            ["项目编号", "零件号", "零件描述", "工厂", "立项日期"],
            ["CN-1", "P1", "Housing", "Changchun", date(2026, 1, 2)],
        ],
        title="Master Changes",
    )
    response = await client.post("/api/excel/import", files=upload(content))
    assert response.status_code == 200, response.text
    project = (await client.get("/api/projects/CN-1")).json()
    assert project["created_at"] == "2026-01-02"
    assert project["plant"] == "Changchun"


async def test_invalid_row_rolls_back_entire_workbook(client):
    content = workbook(
        [["id", "pn", "desc"], ["XL-1", "P1", "Good"], ["XL-2", None, "Missing part"]]
    )
    preview = (await client.post("/api/excel/preview", files=upload(content))).json()
    assert not preview["valid"]
    assert preview["errors"][0]["row"] == 3
    assert (
        await client.post("/api/excel/import", files=upload(content))
    ).status_code == 422
    assert (await client.get("/api/projects")).json() == []


async def test_create_conflict_and_explicit_upsert(client):
    await client.post(
        "/api/projects",
        json={"id": "XL-1", "pn": "old", "desc": "Original", "plant": "Keep"},
    )
    content = workbook([["id", "pn", "desc"], ["XL-1", "new", "Updated"]])
    assert not (await client.post("/api/excel/preview", files=upload(content))).json()[
        "valid"
    ]
    assert (
        await client.post("/api/excel/import", files=upload(content))
    ).status_code == 422
    assert (
        await client.post("/api/excel/import?mode=upsert", files=upload(content))
    ).status_code == 200
    project = (await client.get("/api/projects/XL-1")).json()
    assert project["pn"] == "new"
    assert project["plant"] == "Keep"


async def test_export_reimport_roundtrip_preserves_values_and_links(client, db_engine):
    await client.post(
        "/api/projects",
        json={
            "id": "RT-1",
            "pn": "00012",
            "desc": "Roundtrip",
            "created_at": "2026-01-01",
            "team": {"SDE": "Alice"},
        },
    )
    await client.patch(
        "/api/projects/RT-1/actions/RT-1-0",
        json={
            "value": "=not_a_formula",
            "status": "approved",
            "comment": "note",
            "link": "https://example.com/file",
            "done_date": "2026-01-02",
        },
    )
    await client.post(
        "/api/projects/RT-1/actions",
        json={
            "ph": 0,
            "tab": "Review",
            "act": "Review",
            "input_type": "Date",
            "fn": "SDE",
            "lead": None,
            "due_date": "2026-01-10",
        },
    )
    for value in ("2026-04-01", "2026-11-01"):
        response = await client.patch(
            "/api/projects/RT-1/actions/RT-1-2", json={"value": value}
        )
        assert response.status_code == 200
    before = (await client.get("/api/projects/RT-1")).json()
    response = await client.get("/api/excel/export?project_id=RT-1")
    assert response.status_code == 200
    content = response.content
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ["Projects", "Teams", "Actions", "Summary"]
    text_cells = [
        cell
        for row in book["Actions"]
        for cell in row
        if cell.value == "=not_a_formula"
    ]
    assert len(text_cells) == 1 and text_cells[0].data_type == "s"
    book.close()
    async with async_sessionmaker(db_engine)() as session:
        await session.execute(delete(Project))
        await session.commit()
    response = await client.post("/api/excel/import", files=upload(content))
    assert response.status_code == 200, response.text
    after = (await client.get("/api/projects/RT-1")).json()
    for key in ("id", "pn", "desc", "created_at", "actions", "summary"):
        assert after[key] == before[key], key
    assert [(t["function"], t["name"]) for t in after["team_members"]] == [
        (t["function"], t["name"]) for t in before["team_members"]
    ]


@pytest.mark.parametrize("kind", ["duplicate", "cycle", "cross_project"])
async def test_invalid_graph_or_ids_reject_workbook(client, kind):
    rows = [["id", "pn", "desc"], ["XL-1", "P1", "Test"]]
    action_rows = [
        [
            "id",
            "project_id",
            "ph",
            "tab",
            "act",
            "input_type",
            "fn",
            "dep_id",
            "is_custom",
            "lead",
        ],
        ["custom-a", "XL-1", 1, "A", "A", "Date", "SDE", "custom-b", True, 1],
        ["custom-b", "XL-1", 1, "B", "B", "Date", "SDE", "custom-a", True, 1],
    ]
    if kind == "duplicate":
        rows.append(["XL-1", "P2", "Duplicate"])
        action_rows = None
    if kind == "cross_project":
        assert action_rows is not None
        action_rows[1][7] = "other-project-action"
    content = workbook(rows, actions=action_rows)
    assert (
        await client.post("/api/excel/import", files=upload(content))
    ).status_code == 422
    assert (await client.get("/api/projects")).json() == []


async def test_bad_file_and_formula_report_errors(client):
    assert (
        await client.post("/api/excel/preview", files=upload(b"not excel"))
    ).status_code == 422
    content = workbook([["id", "pn", "desc"], ["XL-1", "P1", '=HYPERLINK("bad")']])
    report = (await client.post("/api/excel/preview", files=upload(content))).json()
    assert not report["valid"]
    assert "Formula" in report["errors"][0]["message"]


async def test_template_and_actions_filter(client):
    template = await client.get("/api/excel/template")
    report = (
        await client.post("/api/excel/preview", files=upload(template.content))
    ).json()
    assert report["valid"] and report["created_projects"] == 0
    await client.post(
        "/api/projects",
        json={
            "id": "LIST-1",
            "pn": "P",
            "desc": "Action list",
            "team": {"SDE": "Alice"},
        },
    )
    response = await client.get("/api/actions?owner=Alice")
    assert response.status_code == 200
    assert all(
        a["owner"] == "Alice" and a["project_description"] == "Action list"
        for a in response.json()
    )


async def test_corrupt_worksheet_returns_validation_error(client):
    from zipfile import ZipFile

    content = workbook([["id", "pn", "desc"], ["XL-1", "P1", "Test"]])
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w") as destination:
        for item in source.infolist():
            destination.writestr(
                item,
                b"<invalid"
                if item.filename == "xl/worksheets/sheet1.xml"
                else source.read(item.filename),
            )
    response = await client.post("/api/excel/preview", files=upload(output.getvalue()))
    assert response.status_code == 422
    assert (await client.get("/api/projects")).json() == []


async def test_imported_date_changes_preserve_baseline_and_audit_history(client):
    await client.post("/api/projects", json={"id": "H-1", "pn": "P", "desc": "History"})
    url = "/api/projects/H-1/actions/H-1-2"
    await client.patch(url, json={"value": "2026-04-01", "done_date": "2026-09-01"})
    rule = STANDARD_MATRIX[2]

    def content(value):
        return workbook(
            [["id", "pn", "desc"], ["H-1", "P", "History"]],
            actions=[
                [
                    "id",
                    "project_id",
                    "ph",
                    "tab",
                    "act",
                    "input_type",
                    "fn",
                    "value",
                    "orig",
                    "date_log",
                ],
                [
                    "H-1-2",
                    "H-1",
                    rule["ph"],
                    rule["tab"],
                    rule["act"],
                    rule["input_type"],
                    rule["fn"],
                    value,
                    "2026-11-01",
                    "[]",
                ],
            ],
        )

    before = (await client.get("/api/projects/H-1")).json()
    response = await client.post(
        "/api/excel/preview?mode=upsert", files=upload(content("2026-11-01"))
    )
    assert response.status_code == 200, response.text
    assert (await client.get("/api/projects/H-1")).json() == before
    for value in ("2026-11-01", "2026-11-01", "", "2026-12-01"):
        response = await client.post(
            "/api/excel/import?mode=upsert", files=upload(content(value))
        )
        assert response.status_code == 200, response.text
    project = (await client.get("/api/projects/H-1")).json()
    action = next(a for a in project["actions"] if a["id"] == "H-1-2")
    assert action["orig"] == "2026-04-01"
    assert action["value"] == "2026-12-01"
    assert action["done_date"] is None
    assert [(e["from"], e["to"], e["source"]) for e in action["date_log"]] == [
        ("", "2026-04-01", "edit"),
        ("2026-04-01", "2026-11-01", "import"),
        ("2026-11-01", "", "import"),
        ("", "2026-12-01", "import"),
    ]
    response = await client.post(
        "/api/excel/import?mode=upsert", files=upload(content("2026-02-30"))
    )
    assert response.status_code == 422
    assert (await client.get("/api/projects/H-1")).json() == project


@pytest.mark.parametrize("defect", ["gap", "chronology", "current"])
async def test_fresh_import_rejects_inconsistent_date_history(client, defect):
    history = [
        {
            "ts": "2026-09-21T01:00:00Z",
            "by": "Alice",
            "from": "",
            "to": "2026-04-01",
            "source": "edit",
        },
        {
            "ts": "2026-09-22T01:00:00Z",
            "by": "Alice",
            "from": "2026-04-01",
            "to": "2026-11-01",
            "source": "edit",
        },
    ]
    if defect == "gap":
        history[1]["from"] = "2026-05-01"
    elif defect == "chronology":
        history[1]["ts"] = "2026-09-20T01:00:00Z"
    else:
        history[1]["to"] = "2026-12-01"
    rule = STANDARD_MATRIX[2]
    content = workbook(
        [["id", "pn", "desc"], ["H-1", "P", "History"]],
        actions=[
            [
                "id",
                "project_id",
                "ph",
                "tab",
                "act",
                "input_type",
                "fn",
                "value",
                "date_log",
            ],
            [
                "H-1-2",
                "H-1",
                rule["ph"],
                rule["tab"],
                rule["act"],
                rule["input_type"],
                rule["fn"],
                "2026-11-01",
                json.dumps(history),
            ],
        ],
    )
    response = await client.post("/api/excel/import", files=upload(content))
    assert response.status_code == 422, response.text
    assert "Date history" in response.text
    assert (await client.get("/api/projects")).json() == []


async def test_custom_import_cannot_confirm_placeholder_reference(client):
    content = workbook(
        [["id", "pn", "desc"], ["H-1", "P", "History"]],
        actions=[
            [
                "id",
                "project_id",
                "ph",
                "tab",
                "act",
                "input_type",
                "fn",
                "value",
                "status",
                "done_date",
                "is_custom",
            ],
            [
                "custom-1",
                "H-1",
                1,
                "Check",
                "Check",
                "Reference",
                "SDE",
                "TBD",
                "approved",
                "2026-09-01",
                True,
            ],
        ],
    )
    response = await client.post("/api/excel/import", files=upload(content))
    assert response.status_code == 422, response.text
    assert (await client.get("/api/projects")).json() == []


@pytest.mark.parametrize(
    ("index", "value", "extra_headers", "extra_values", "completed"),
    [
        (1, "PPAP-1", [], [], False),
        (1, "PPAP-1", ["done_date"], ["2026-09-01"], True),
        (0, "CR-1", ["done_date"], ["2026-09-01"], False),
        (0, "CR-1", ["status", "done_date"], [None, "2026-09-01"], False),
        (0, "CR-1", ["status", "done_date"], ["in_progress", "2026-09-01"], False),
        (0, "CR-1", ["status"], ["approved"], True),
        (0, "CR-1", ["status", "done_date"], ["approved", None], True),
        (11, "Draft", ["done_date"], ["2026-09-01"], False),
        (11, "Full approved", [], [], True),
        (11, "Full approved", ["done_date"], [None], True),
    ],
)
async def test_reimport_reassesses_unchanged_completion_from_imported_evidence(
    client, db_engine, index, value, extra_headers, extra_values, completed
):
    await client.post(
        "/api/projects", json={"id": "R-1", "pn": "P", "desc": "Reimport"}
    )
    async with db_engine.begin() as connection:
        await connection.execute(
            text(
                "UPDATE project_actions SET value=:value,done_date='2026-09-01' WHERE id=:id"
            ),
            {"value": value, "id": f"R-1-{index}"},
        )
    rule = STANDARD_MATRIX[index]
    content = workbook(
        [["id", "pn", "desc"], ["R-1", "P", "Reimport"]],
        actions=[
            [
                "id",
                "project_id",
                "ph",
                "tab",
                "act",
                "input_type",
                "fn",
                "value",
                *extra_headers,
            ],
            [
                f"R-1-{index}",
                "R-1",
                rule["ph"],
                rule["tab"],
                rule["act"],
                rule["input_type"],
                rule["fn"],
                value,
                *extra_values,
            ],
        ],
    )
    original = (await client.get("/api/projects/R-1")).json()
    response = await client.post(
        "/api/excel/preview?mode=upsert", files=upload(content)
    )
    assert response.status_code == 200, response.text
    assert response.json()["valid"], response.text
    if not completed:
        assert any("reopened" in warning for warning in response.json()["warnings"])
    assert (await client.get("/api/projects/R-1")).json() == original
    response = await client.post("/api/excel/import?mode=upsert", files=upload(content))
    assert response.status_code == 200, response.text
    result = (await client.get("/api/projects/R-1")).json()
    action = next(a for a in result["actions"] if a["id"] == f"R-1-{index}")
    assert bool(action["done_date"]) == completed
    assert action["value"] == value
    assert action["date_log"] == []
    current_actions = {a["id"]: a for a in result["actions"]}
    for previous in original["actions"]:
        if previous["id"] != action["id"]:
            for key in (
                "id",
                "value",
                "done_date",
                "approval_status",
                "orig",
                "date_log",
                "dep_id",
            ):
                assert current_actions[previous["id"]][key] == previous[key]


async def test_reimport_does_not_inherit_old_approval_when_current_file_omits_it(
    client,
):
    await client.post(
        "/api/projects", json={"id": "R-1", "pn": "P", "desc": "Reimport"}
    )
    await client.patch(
        "/api/projects/R-1/actions/R-1-0",
        json={"value": "CR-1", "status": "approved", "done_date": "2026-09-01"},
    )
    rule = STANDARD_MATRIX[0]
    content = workbook(
        [["id", "pn", "desc"], ["R-1", "P", "Reimport"]],
        actions=[
            [
                "id",
                "project_id",
                "ph",
                "tab",
                "act",
                "input_type",
                "fn",
                "value",
                "done_date",
            ],
            [
                "R-1-0",
                "R-1",
                rule["ph"],
                rule["tab"],
                rule["act"],
                rule["input_type"],
                rule["fn"],
                "CR-1",
                "2026-09-01",
            ],
        ],
    )
    response = await client.post("/api/excel/import?mode=upsert", files=upload(content))
    assert response.status_code == 200, response.text
    assert any("reopened" in warning for warning in response.json()["warnings"])
    action = (await client.get("/api/projects/R-1")).json()["actions"][0]
    assert action["done_date"] is None
    assert action["approval_status"] == "in_progress"


async def test_reimport_of_only_comments_keeps_unmentioned_completion(client):
    await client.post(
        "/api/projects", json={"id": "R-1", "pn": "P", "desc": "Reimport"}
    )
    await client.patch(
        "/api/projects/R-1/actions/R-1-0",
        json={"value": "CR-1", "status": "approved", "done_date": "2026-09-01"},
    )
    rule = STANDARD_MATRIX[0]
    content = workbook(
        [["id", "pn", "desc"], ["R-1", "P", "Reimport"]],
        actions=[
            ["id", "project_id", "ph", "tab", "act", "input_type", "fn", "comment"],
            [
                "R-1-0",
                "R-1",
                rule["ph"],
                rule["tab"],
                rule["act"],
                rule["input_type"],
                rule["fn"],
                "Updated comment",
            ],
        ],
    )
    response = await client.post("/api/excel/import?mode=upsert", files=upload(content))
    assert response.status_code == 200, response.text
    action = (await client.get("/api/projects/R-1")).json()["actions"][0]
    assert action["comment"] == "Updated comment"
    assert action["done_date"] == "2026-09-01"
    assert action["approval_status"] == "approved"


async def test_reimport_reopens_unchanged_date_without_fabricating_history(client):
    await client.post(
        "/api/projects", json={"id": "R-1", "pn": "P", "desc": "Reimport"}
    )
    url = "/api/projects/R-1/actions/R-1-2"
    await client.patch(url, json={"value": "2026-04-01"})
    response = await client.patch(
        url, json={"value": "2026-11-01", "done_date": "2026-09-01"}
    )
    original = response.json()
    rule = STANDARD_MATRIX[2]
    content = workbook(
        [["id", "pn", "desc"], ["R-1", "P", "Reimport"]],
        actions=[
            [
                "id",
                "project_id",
                "ph",
                "tab",
                "act",
                "input_type",
                "fn",
                "value",
                "orig",
                "date_log",
            ],
            [
                "R-1-2",
                "R-1",
                rule["ph"],
                rule["tab"],
                rule["act"],
                rule["input_type"],
                rule["fn"],
                "2026-11-01",
                "2026-11-01",
                "[]",
            ],
        ],
    )
    for _ in range(2):
        response = await client.post(
            "/api/excel/import?mode=upsert", files=upload(content)
        )
        assert response.status_code == 200, response.text
        action = next(
            a
            for a in (await client.get("/api/projects/R-1")).json()["actions"]
            if a["id"] == "R-1-2"
        )
        assert action["done_date"] is None
        for key in ("orig", "date_log", "value", "dep_id", "id"):
            assert action[key] == original[key]


@pytest.mark.parametrize(
    ("index", "value", "status"), [(0, "CR-1", "approved"), (11, "Full approved", None)]
)
async def test_reimport_same_final_approval_preserves_actual_completion_date(
    client, index, value, status
):
    await client.post(
        "/api/projects", json={"id": "R-1", "pn": "P", "desc": "Reimport"}
    )
    response = await client.patch(
        f"/api/projects/R-1/actions/R-1-{index}",
        json={"value": value, "status": status, "done_date": "2026-09-01"},
    )
    assert response.status_code == 200, response.text
    rule = STANDARD_MATRIX[index]
    content = workbook(
        [["id", "pn", "desc"], ["R-1", "P", "Reimport"]],
        actions=[
            [
                "id",
                "project_id",
                "ph",
                "tab",
                "act",
                "input_type",
                "fn",
                "value",
                "status",
                "done_date",
            ],
            [
                f"R-1-{index}",
                "R-1",
                rule["ph"],
                rule["tab"],
                rule["act"],
                rule["input_type"],
                rule["fn"],
                value,
                status,
                None,
            ],
        ],
    )
    for _ in range(2):
        response = await client.post(
            "/api/excel/import?mode=upsert", files=upload(content)
        )
        assert response.status_code == 200, response.text
        action = next(
            a
            for a in (await client.get("/api/projects/R-1")).json()["actions"]
            if a["id"] == f"R-1-{index}"
        )
        assert action["done_date"] == "2026-09-01"
